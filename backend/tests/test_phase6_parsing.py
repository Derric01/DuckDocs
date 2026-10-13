"""Real synthetic Office packages; never open a scorecard document."""
import io

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement

from app.core.config import Settings
from app.services.chunking import chunk_document
from app.services.parsing import parse_docx


def payload(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def parsed_chunks(document, **settings):
    parsed = parse_docx(payload(document))
    assert parsed is not None
    return parsed, chunk_document(parsed, Settings(**settings))


def test_explicit_and_rendered_break_in_same_run_count_once():
    document = Document()
    run = document.add_paragraph().add_run("Cedar opening text.")
    run.add_break(WD_BREAK.PAGE)
    run._r.append(OxmlElement("w:lastRenderedPageBreak"))
    run.add_text("Birch continuation text.")
    document.add_paragraph("Maple next paragraph.")
    parsed, chunks = parsed_chunks(document)
    assert parsed.page_count == 2
    assert [(c.page, c.paragraph_number, c.text) for c in chunks] == [
        (1, 1, "Cedar opening text."), (2, 1, "Birch continuation text."), (2, 2, "Maple next paragraph.")]


def test_rendered_break_without_explicit_break_splits_a_paragraph():
    document = Document()
    run = document.add_paragraph().add_run("Opening statement.")
    run._r.append(OxmlElement("w:lastRenderedPageBreak"))
    run.add_text("Continuation statement.")
    parsed, chunks = parsed_chunks(document)
    assert parsed.page_count == 2
    assert [(c.page, c.paragraph_number) for c in chunks] == [(1, 1), (2, 1)]


def test_page_break_before_inherited_style_and_explicit_false_override():
    document = Document()
    document.styles["Heading 1"].paragraph_format.page_break_before = True
    document.add_paragraph("Opening page.")
    document.add_paragraph("Chapter title.", "Heading 1")
    paragraph = document.add_paragraph("Same page heading.", "Heading 1")
    paragraph.paragraph_format.page_break_before = False
    paragraph = document.add_paragraph("Next page.")
    paragraph.paragraph_format.page_break_before = True
    paragraph.runs[0]._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))
    parsed, chunks = parsed_chunks(document)
    assert parsed.page_count == 3
    assert [c.page for c in chunks] == [1, 2, 2, 3]


@pytest.mark.parametrize("start,expected", [
    (WD_SECTION_START.NEW_PAGE, 2), (WD_SECTION_START.CONTINUOUS, 1),
    (WD_SECTION_START.EVEN_PAGE, 2), (WD_SECTION_START.ODD_PAGE, 3),
])
def test_section_start_type_belongs_to_the_following_section(start, expected):
    document = Document()
    document.add_paragraph("First section text.")
    document.add_section(start)
    document.add_paragraph("Following section text.")
    parsed, chunks = parsed_chunks(document)
    assert parsed.page_count == expected
    assert [(c.page, c.text) for c in chunks] == [(1, "First section text."), (expected, "Following section text.")]


def test_empty_pages_keep_physical_numbering_including_a_trailing_page():
    document = Document()
    document.add_paragraph("Opening text.")
    document.add_page_break()
    document.add_page_break()
    document.add_paragraph("Third page text.")
    document.add_page_break()
    parsed, chunks = parsed_chunks(document)
    assert parsed.page_count == 4
    assert [c.page for c in chunks] == [1, 3]


def test_cells_and_body_paragraphs_preserve_pages_and_reading_order():
    document = Document()
    document.add_paragraph("Before table.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left cell."
    run = table.cell(0, 1).paragraphs[0].add_run("Right cell before break.")
    run.add_break(WD_BREAK.PAGE)
    run.add_text("Right cell after break.")
    document.add_paragraph("After table.")
    parsed, chunks = parsed_chunks(document)
    assert parsed.page_count == 2
    assert [c.page for c in chunks] == [1, 1, 1, 2, 2]
    assert [c.paragraph_number for c in chunks] == [1, None, None, None, 1]
    assert chunks[1].anchor_quality == chunks[3].anchor_quality == "cell"
    assert chunks[3].section == "Table 1 / Row 1 / Cell 2"


def test_cells_share_row_start_while_following_rows_start_after_spanning_cell():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    run = table.cell(0, 0).paragraphs[0].add_run("Left cell opening.")
    run.add_break(WD_BREAK.PAGE)
    run.add_text("Left cell continuation.")
    table.cell(0, 1).text = "Right cell still starts on the first page."
    table.cell(1, 0).text = "Following row starts on the second page."
    document.add_paragraph("Body after the table.")
    parsed, chunks = parsed_chunks(document)
    assert parsed.page_count == 2
    assert {c.text: c.page for c in chunks} == {
        "Left cell opening.": 1, "Left cell continuation.": 2,
        "Right cell still starts on the first page.": 1,
        "Following row starts on the second page.": 2, "Body after the table.": 2,
    }


def test_headers_and_footers_are_repeated_on_correct_pages_without_body_ordinals():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Running header."
    document.sections[0].footer.paragraphs[0].text = "Running footer."
    document.add_paragraph("Body opening.")
    document.add_page_break()
    document.add_paragraph("Body ending.")
    _, chunks = parsed_chunks(document)
    assert [(c.page, c.paragraph_number) for c in chunks if c.section.startswith("Paragraph")] == [(1, 1), (2, 1)]
    for kind in ("Header", "Footer"):
        assert [(c.page, c.paragraph_number) for c in chunks if c.section == kind] == [(1, None), (2, None)]


def test_long_paragraph_windows_keep_one_ordinal_and_do_not_merge_neighbors():
    document = Document()
    document.add_paragraph(" ".join(f"token{i}" for i in range(60)))
    document.add_paragraph("Second short paragraph.")
    _, chunks = parsed_chunks(document, chunk_max_words=25, chunk_overlap_words=5)
    assert [c.paragraph_number for c in chunks] == [1, 1, 1, 2]
    assert chunks[0].text.split()[-5:] == chunks[1].text.split()[:5]


@pytest.mark.parametrize("storage", ["memory", "sql"])
def test_docx_pages_and_ordinals_survive_repository_restart(tmp_path, storage):
    from app.domain.models import Document as StoredDocument, utc_now
    from app.repositories.memory import DocumentRepository
    from app.repositories.sql import DocumentRow, SqlDocumentRepository
    factory = (lambda: DocumentRepository(tmp_path)) if storage == "memory" else (
        lambda: SqlDocumentRepository(f"sqlite:///{tmp_path / 'paragraphs.db'}"))
    repo = factory()
    now = utc_now()
    stored = StoredDocument(id="doc_pages", name="report.docx", file_type="docx", mime_type="application/octet-stream",
                            size_bytes=20, status="ready", fidelity_tier="structural", pages=2,
                            category="Document", current_version_id="v1", created_at=now, updated_at=now)
    repo.documents[stored.id] = stored
    if storage == "sql":
        with repo.SessionLocal() as session:
            session.merge(DocumentRow(**stored.model_dump()))
            session.commit()
    document = Document()
    document.add_paragraph("Opening body.")
    document.add_page_break()
    document.add_paragraph("Closing body.")
    _, chunks = parsed_chunks(document)
    repo.index_document_chunks(stored.id, chunks)
    assert sorted((e.page, e.paragraph_number) for e in factory().all_evidence()) == [(1, 1), (2, 1)]


def test_migration_keeps_legacy_evidence_and_unknown_paragraph_null(tmp_path):
    from alembic import command
    from alembic.config import Config
    from sqlalchemy import create_engine, text
    from pathlib import Path
    backend = Path(__file__).resolve().parents[1]
    config = Config(str(backend / "alembic.ini"))
    config.set_main_option("script_location", str(backend / "migrations"))
    url = f"sqlite:///{tmp_path / 'upgrade.db'}"
    config.set_main_option("sqlalchemy.url", url)
    engine = create_engine(url)
    command.upgrade(config, "0005")
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO evidence (id, document_id, document_name, section, page, line_start, line_end, snippet, retrieval_score, relevance) VALUES ('old', 'doc', 'notes.docx', 'Imported', 7, 1, 1, 'Legacy text.', 0.8, 'High')"))
    command.upgrade(config, "head")
    with engine.connect() as connection:
        assert connection.execute(text("SELECT page, paragraph_number, snippet FROM evidence")).one() == (7, None, "Legacy text.")
