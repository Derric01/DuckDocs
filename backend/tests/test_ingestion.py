"""Ingestion pipeline tests: real PDF parsing, unlimited local OCR, and
page-aware chunking (docs/21_FILE_PROCESSING.md, docs/22_OCR_PIPELINE.md).

These exercise the actual Tesseract/PyMuPDF code paths (no mocking of OCR)
so a regression that silently breaks recognition -- e.g. reintroducing the
old "PDFs/images never get their text extracted" bug -- fails loudly.
"""

from __future__ import annotations

import io
from pathlib import Path

import fitz
import openpyxl
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont

from app.core.config import Settings
from app.main import app, get_registry, get_repository, get_runtime_settings, get_vector_store
from app.providers.registry import ProviderConfig, ProviderRegistry
from app.repositories.memory import DocumentRepository
from app.services.chunking import ChunkCandidate, chunk_document
from app.services.parsing import (
    ParsedDocument,
    ParsedPage,
    parse_csv,
    parse_html,
    parse_image,
    parse_pdf,
    parse_upload,
    parse_xlsx,
)
from app.services.vector_store import VectorStore

SETTINGS = Settings.from_env()


def _render_text_image(text: str, size: tuple[int, int] = (900, 220)) -> Image.Image:
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 40)
    except OSError:
        font = ImageFont.load_default()
    draw.text((30, 80), text, fill="black", font=font)
    return image


def _image_bytes(text: str, fmt: str = "PNG") -> bytes:
    buffer = io.BytesIO()
    _render_text_image(text).save(buffer, format=fmt)
    return buffer.getvalue()


def _scanned_pdf_bytes(text: str) -> bytes:
    """A PDF with NO text layer at all -- only a rasterized page image."""
    with io.BytesIO() as png_buffer:
        _render_text_image(text).save(png_buffer, format="PNG")
        png_bytes = png_buffer.getvalue()
    document = fitz.open()
    page = document.new_page(width=900, height=220)
    page.insert_image(page.rect, stream=png_bytes)
    return document.tobytes()


def _native_pdf_bytes(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text, fontsize=14)
    return document.tobytes()


def client_with_data_root(data_root: Path) -> TestClient:
    repo = DocumentRepository(data_root)
    runtime = Settings(data_root=data_root, ollama_base_url="http://127.0.0.1:9", chroma_url=None, db_url=None)
    registry = ProviderRegistry(runtime, config_path=data_root / "provider_configs.json")
    registry.upsert_config(
        ProviderConfig(
            id="cfg_chat_extractive", role="chat", provider_type="extractive", model_name="evidence_synthesis_v1", is_default=True
        )
    )
    registry.upsert_config(
        ProviderConfig(
            id="cfg_embed_keyword", role="embedding", provider_type="keyword", model_name="local_text_match", is_default=True
        )
    )
    store = VectorStore(runtime)
    app.dependency_overrides[get_repository] = lambda: repo
    app.dependency_overrides[get_runtime_settings] = lambda: runtime
    app.dependency_overrides[get_registry] = lambda: registry
    app.dependency_overrides[get_vector_store] = lambda: store
    return TestClient(app)


# ---------------------------------------------------------------------------
# Unit tests: parsing
# ---------------------------------------------------------------------------


def test_native_pdf_extracts_text_without_ocr() -> None:
    payload = _native_pdf_bytes("Quarterly compliance review is due every March.")
    parsed = parse_pdf(payload, SETTINGS)
    assert parsed is not None
    assert parsed.fidelity_tier == "full_layout"
    assert parsed.pages[0].source == "native"
    assert parsed.pages[0].confidence is None
    assert "compliance review" in parsed.pages[0].text


def test_scanned_pdf_with_no_text_layer_is_ocrd() -> None:
    payload = _scanned_pdf_bytes("Retention window is eighteen months")
    events: list[tuple[int, int]] = []
    parsed = parse_pdf(payload, SETTINGS, on_page=lambda current, total: events.append((current, total)))
    assert parsed is not None
    assert parsed.fidelity_tier == "ocr_dependent"
    assert parsed.pages[0].source == "ocr"
    assert parsed.pages[0].confidence is not None
    assert 0.0 <= parsed.pages[0].confidence <= 1.0
    assert "Retention" in parsed.pages[0].text or "retention" in parsed.pages[0].text.lower()
    assert events == [(1, 1)]


def test_every_scanned_page_is_ocrd_not_just_a_prefix() -> None:
    """No artificial page cap: a multi-page scanned PDF gets OCR on every page."""
    document = fitz.open()
    page_count = 6
    for index in range(page_count):
        page = document.new_page(width=900, height=220)
        image_bytes = _image_bytes(f"Scanned page number {index + 1}")
        page.insert_image(page.rect, stream=image_bytes)
    payload = document.tobytes()

    events: list[tuple[int, int]] = []
    parsed = parse_pdf(payload, SETTINGS, on_page=lambda current, total: events.append((current, total)))
    assert parsed is not None
    assert len(parsed.pages) == page_count
    assert all(page.source == "ocr" for page in parsed.pages)
    assert events[-1] == (page_count, page_count)
    assert any(page.text.strip() for page in parsed.pages)


def test_image_upload_is_ocrd() -> None:
    payload = _image_bytes("Falcon audit trail 2026")
    parsed = parse_image(payload, SETTINGS)
    assert parsed is not None
    assert parsed.fidelity_tier == "ocr_dependent"
    assert parsed.pages[0].source == "ocr"
    assert "Falcon" in parsed.pages[0].text or "audit" in parsed.pages[0].text.lower()


def test_garbage_pdf_bytes_fail_closed_without_crashing() -> None:
    assert parse_pdf(b"not a real pdf", SETTINGS) is None


def test_xlsx_parses_sheet_rows_with_structural_tier() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["Item", "Cost"])
    sheet.append(["Widgets", 42])
    buffer = io.BytesIO()
    workbook.save(buffer)
    parsed = parse_xlsx(buffer.getvalue())
    assert parsed is not None
    assert parsed.fidelity_tier == "structural"
    assert "Widgets" in parsed.pages[0].text


def test_csv_and_html_parse_with_structural_tier() -> None:
    csv_parsed = parse_csv(b"name,score\nAlice,95\nBob,88\n")
    assert csv_parsed is not None
    assert csv_parsed.fidelity_tier == "structural"
    assert "Alice" in csv_parsed.pages[0].text

    html_parsed = parse_html(b"<html><body><h1>Report</h1><p>Renewals happen every quarter.</p></body></html>")
    assert html_parsed is not None
    assert "Renewals happen every quarter" in html_parsed.pages[0].text


def test_unsupported_extension_returns_none() -> None:
    assert parse_upload(b"binary junk", "exe", SETTINGS) is None


# ---------------------------------------------------------------------------
# Unit tests: chunking
# ---------------------------------------------------------------------------


def test_chunking_preserves_real_page_numbers_and_ocr_confidence() -> None:
    parsed = ParsedDocument(
        pages=[
            ParsedPage(1, "Native page one.\n\nSecond paragraph here.", "native", None),
            ParsedPage(2, "OCR recognized text on page two.", "ocr", 0.42),
        ],
        fidelity_tier="ocr_dependent",
        page_count=2,
    )
    chunks = chunk_document(parsed, SETTINGS)
    assert [c.page for c in chunks] == [1, 2]
    assert chunks[0].ocr_confidence is None
    assert chunks[1].ocr_confidence == 0.42
    assert chunks[1].fidelity_tier == "ocr_dependent"


def test_chunking_splits_long_pages_on_word_budget_with_overlap() -> None:
    settings = Settings(chunk_max_words=20, chunk_overlap_words=5)
    paragraphs = [f"Paragraph {i} has some words in it for budget testing." for i in range(10)]
    page = ParsedPage(1, "\n\n".join(paragraphs), "native", None)
    parsed = ParsedDocument(pages=[page], fidelity_tier="full_layout", page_count=1)
    chunks = chunk_document(parsed, settings)
    assert len(chunks) > 1
    assert all(c.page == 1 for c in chunks)
    # Consecutive chunks share at least one paragraph's worth of overlap text.
    assert chunks[0].text.split()[-1] in chunks[1].text.split() or chunks[0].line_end >= chunks[1].line_start


def test_chunk_document_respects_max_chunks_ceiling() -> None:
    settings = Settings(max_chunks_per_document=3, chunk_max_words=5)
    text = "\n\n".join(f"word{i} filler text here" for i in range(50))
    parsed = ParsedDocument(pages=[ParsedPage(1, text, "native", None)], fidelity_tier="full_layout", page_count=1)
    chunks = chunk_document(parsed, settings)
    assert len(chunks) == 3


def test_chunk_candidate_is_plain_dataclass() -> None:
    candidate = ChunkCandidate(
        page=1, line_start=1, line_end=1, text="hello", ocr_confidence=None, anchor_quality="paragraph", fidelity_tier="full_layout"
    )
    assert candidate.text == "hello"


# ---------------------------------------------------------------------------
# API-level: full upload -> OCR -> searchable evidence
# ---------------------------------------------------------------------------


def test_scanned_pdf_upload_becomes_searchable_with_ocr_confidence(tmp_path: Path) -> None:
    with client_with_data_root(tmp_path) as client:
        payload = _scanned_pdf_bytes("Northwind lease renews every eighteen months")
        upload = client.post(
            "/api/v1/documents",
            files={"files": ("scanned-lease.pdf", payload, "application/pdf")},
        )
        assert upload.status_code == 201
        document_id = upload.json()["items"][0]["document"]["id"]

        document = client.get(f"/api/v1/documents/{document_id}").json()
        assert document["status"] == "ready"
        assert document["fidelity_tier"] == "ocr_dependent"

        search = client.post("/api/v1/search", json={"query": "northwind lease renews"}).json()
        assert search["results"], "scanned PDF with no text layer must still be searchable via OCR"
        evidence = search["results"][0]["evidence"]
        assert evidence["fidelity_tier"] == "ocr_dependent"
        assert evidence["ocr_confidence"] is not None
        assert 0.0 <= evidence["ocr_confidence"] <= 1.0

    app.dependency_overrides.clear()


def test_image_upload_becomes_searchable(tmp_path: Path) -> None:
    with client_with_data_root(tmp_path) as client:
        payload = _image_bytes("Amberstone obligations renew quarterly")
        upload = client.post("/api/v1/documents", files={"files": ("note.png", payload, "image/png")})
        assert upload.status_code == 201
        document_id = upload.json()["items"][0]["document"]["id"]

        document = client.get(f"/api/v1/documents/{document_id}").json()
        assert document["fidelity_tier"] == "ocr_dependent"

        search = client.post("/api/v1/search", json={"query": "amberstone obligations renew"}).json()
        assert search["results"]

    app.dependency_overrides.clear()


def test_unreadable_file_marks_document_for_review_not_ready(tmp_path: Path) -> None:
    with client_with_data_root(tmp_path) as client:
        # A .txt file with no alphanumeric content: fails the extraction floor.
        upload = client.post(
            "/api/v1/documents",
            files={"files": ("empty.txt", b"....... ---- .......", "text/plain")},
        )
        assert upload.status_code == 201
        document_id = upload.json()["items"][0]["document"]["id"]
        document = client.get(f"/api/v1/documents/{document_id}").json()
        assert document["status"] == "review"

    app.dependency_overrides.clear()


def test_unsupported_extension_is_rejected_at_upload(tmp_path: Path) -> None:
    with client_with_data_root(tmp_path) as client:
        upload = client.post(
            "/api/v1/documents",
            files={"files": ("payload.exe", b"MZ\x90\x00", "application/octet-stream")},
        )
        assert upload.status_code == 415

    app.dependency_overrides.clear()
