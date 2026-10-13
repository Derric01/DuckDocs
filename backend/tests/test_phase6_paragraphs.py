import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement

from tests.test_ingestion import client_with_data_root
from tests.test_phase5_answers import ask, upload
from tests.test_phase6_parsing import payload


@pytest.mark.parametrize("boundary", ["explicit", "before", "section", "rendered"])
def test_second_paragraph_is_counted_within_destination_page(boundary, tmp_path):
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Running header does not enter body counts."
    document.add_paragraph("Opening text before boundary.")
    if boundary == "explicit":
        document.add_page_break()
    elif boundary == "section":
        document.add_section(WD_SECTION_START.NEW_PAGE)
    first = document.add_heading("Destination heading counts as paragraph one.", level=1)
    if boundary == "before":
        first.paragraph_format.page_break_before = True
    elif boundary == "rendered":
        first.runs[0]._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))
    document.add_paragraph("")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table cell is not a body paragraph."
    expected = "The cedar schedule begins at dawn. A second sentence belongs to the same paragraph."
    document.add_paragraph(expected)
    document.add_paragraph("Third body paragraph is different.")
    document.add_page_break()
    document.add_paragraph("Another page starts counting from one.")
    with client_with_data_root(tmp_path) as client:
        doc = upload(client, "schedule.docx", payload(document))
        result = ask(client, "What does the second paragraph after the page break say?", doc)
        assert result["grounded"]
        assert result["answer"] == expected
        assert result["citations"][0]["page"] == 2
        assert result["citations"][0]["snippet"] == expected
        assert result["diagnostic"]["paragraph_number"] == 2
        assert result["provider"]["name"] == "deterministic"
        assert ask(client, "Quote paragraph 2 on page 2.", doc)["answer"] == expected
        assert ask(client, "Quote paragraph 1 on page 2.", doc)["answer"] == "Destination heading counts as paragraph one."
        missing = ask(client, "Quote the fourth paragraph after the page break.", doc)
        assert not missing["grounded"]
        assert missing["refusal_reason"] == "paragraph_not_found"


def test_mid_paragraph_continuation_counts_as_first_after_break(tmp_path):
    document = Document()
    run = document.add_paragraph().add_run("Before boundary.")
    run.add_break(WD_BREAK.PAGE)
    run.add_text("Continuation is first on the destination page.")
    document.add_paragraph("The next body paragraph is second.")
    with client_with_data_root(tmp_path) as client:
        doc = upload(client, "continuation.docx", payload(document))
        assert ask(client, "Quote the first paragraph after the page break.", doc)["answer"] == "Continuation is first on the destination page."
        assert ask(client, "Quote the second paragraph after the page break.", doc)["answer"] == "The next body paragraph is second."


def test_unspecified_paragraph_is_ambiguous_on_a_page_with_multiple_paragraphs(tmp_path):
    document = Document()
    document.add_paragraph("Opening.")
    document.add_page_break()
    document.add_paragraph("First destination paragraph.")
    document.add_paragraph("Second destination paragraph.")
    with client_with_data_root(tmp_path) as client:
        doc = upload(client, "several.docx", payload(document))
        response = ask(client, "What does the paragraph on the second page say?", doc)
        assert not response["grounded"]
        assert response["refusal_reason"] == "ambiguous_paragraph"


def test_semantic_question_keeps_paragraph_filter_and_normal_grounding(tmp_path):
    document = Document()
    document.add_paragraph("Opening.")
    document.add_page_break()
    document.add_paragraph("The beacon code is AMBER.")
    document.add_paragraph("The beacon code is CEDAR.")
    with client_with_data_root(tmp_path) as client:
        doc = upload(client, "codes.docx", payload(document))
        response = ask(client, "What is the beacon code in the second paragraph on page 2?", doc)
        assert response["grounded"]
        assert "CEDAR" in response["answer"] and "AMBER" not in response["answer"]
        for citation in response["citations"]:
            source = client.get('/api/v1/evidence/' + citation["evidence_unit_id"]).json()
            assert source["paragraph_number"] == source["page"] == 2


def test_split_paragraph_is_not_presented_as_a_complete_quote(tmp_path):
    document = Document()
    document.add_paragraph("Opening.")
    document.add_page_break()
    document.add_paragraph(" ".join(f"term{i}" for i in range(1000)))
    with client_with_data_root(tmp_path) as client:
        doc = upload(client, "lengthy.docx", payload(document))
        response = ask(client, "Quote the first paragraph after the page break.", doc)
        assert not response["grounded"]
        assert response["refusal_reason"] == "paragraph_too_long"


def test_unique_paragraph_on_last_page_is_returned_verbatim(tmp_path):
    document = Document()
    document.add_paragraph("Opening page.")
    document.add_page_break()
    document.add_paragraph("Closing instructions use the northern entrance.")
    with client_with_data_root(tmp_path) as client:
        doc = upload(client, "instructions.docx", payload(document))
        response = ask(client, "What does the paragraph on the last page say?", doc)
        assert response["answer"] == "Closing instructions use the northern entrance."
        assert response["citations"][0]["page"] == 2


def test_unqualified_paragraph_uses_structural_heading_styles_not_text_keywords(tmp_path):
    document = Document()
    document.add_paragraph("Opening page.")
    document.add_page_break()
    document.add_heading("Inventory details", level=1)
    document.add_paragraph("The north depot closes at dusk. Deliveries resume in the morning.")
    document.add_heading("Supporting records", level=2)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Inventory cell is separate."
    with client_with_data_root(tmp_path) as client:
        doc = upload(client, "depot.docx", payload(document))
        result = ask(client, "What does the paragraph on the second page of the Word document say?", doc)
        assert result["answer"] == "The north depot closes at dusk. Deliveries resume in the morning."
        assert result["diagnostic"]["paragraph_number"] == 2
        assert result["citations"][0]["page"] == 2
