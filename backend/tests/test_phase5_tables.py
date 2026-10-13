"""Synthetic structured evidence, with no scorecard/held-out fixture imports."""
import io

import openpyxl
import pytest

from app.core.config import Settings
from app.services.chunking import chunk_document
from app.services.parsing import parse_csv, parse_xlsx


def test_csv_row_coordinates_and_empty_columns_survive_chunking():
    parsed = parse_csv(b'\nPerson,Blank,Amount\nNora,,"1,250"\n\nOmar,x,750\n')
    chunks = chunk_document(parsed, Settings(chunk_max_words=1, chunk_overlap_words=0))
    assert [chunk.line_start for chunk in chunks] == [2, 3, 5]
    assert chunks[1].table_row.cells == ["Nora", "", "1,250"]
    assert chunks[1].table_row.headers == ["Person", "Blank", "Amount"]
    assert chunks[1].table_row.data_row_count == 2
    assert chunks[0].table_row.is_header
    assert chunks[1].text == chunks[1].table_row.quote()
    assert chunks[1].anchor_quality == "cell"


def test_workbook_rows_keep_sheet_and_physical_coordinates():
    workbook = openpyxl.Workbook()
    workbook.active.title = "North"
    workbook.active.append(["Product", "", "Units"])
    workbook.active.append(["Bolt", None, 13])
    workbook.active.append([None, None, None])
    workbook.active.append(["Nut", None, 7])
    workbook.create_sheet("South").append(["Code", "Value"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    chunks = chunk_document(parse_xlsx(buffer.getvalue()), Settings())
    assert [(c.page, c.table_row.sheet_name, c.table_row.row_number) for c in chunks] == [
        (1, "North", 1), (1, "North", 2), (1, "North", 4), (2, "South", 1)]
    assert chunks[2].table_row.cells == ["Nut", "", "7"]
    assert chunks[3].table_row.data_row_count == 0


@pytest.mark.parametrize("storage", ["memory", "sql"])
def test_row_metadata_persists_after_repository_restart(tmp_path, storage):
    from app.domain.models import Document, utc_now
    from app.repositories.memory import DocumentRepository
    from app.repositories.sql import SqlDocumentRepository
    factory = (lambda: DocumentRepository(tmp_path)) if storage == "memory" else (
        lambda: SqlDocumentRepository(f"sqlite:///{tmp_path / 'table.db'}"))
    repo = factory()
    now = utc_now()
    doc = Document(id="doc_rows", name="rows.csv", file_type="csv", mime_type="text/csv",
                   size_bytes=20, status="ready", fidelity_tier="structural", pages=1,
                   category="Data", current_version_id="v1", created_at=now, updated_at=now)
    repo.documents[doc.id] = doc
    if storage == "sql":
        from app.repositories.sql import DocumentRow
        with repo.SessionLocal() as session:
            session.merge(DocumentRow(**doc.model_dump()))
            session.commit()
    chunks = chunk_document(parse_csv(b"Name,Amount\nNora,123\n"), Settings())
    assert repo.index_document_chunks(doc.id, chunks) == 2
    evidence = factory().all_evidence()
    assert len(evidence) == 2
    row = next(e for e in evidence if not e.table_row.is_header)
    assert row.table_row.cells == ["Nora", "123"]
    assert row.file_type == "csv"
    assert row.section == "Row 2"
