from docx import Document

from tests.test_ingestion import client_with_data_root
from tests.test_phase5_answers import ask, upload
from tests.test_phase6_parsing import payload


def test_docx_answer_and_evidence_lookup_agree_on_real_page(tmp_path):
    document = Document()
    document.add_paragraph("The beacon code is AMBER.")
    paragraph = document.add_paragraph("The beacon code is CEDAR.")
    paragraph.paragraph_format.page_break_before = True
    with client_with_data_root(tmp_path) as client:
        doc = upload(client, "guide.docx", payload(document))
        response = ask(client, "What is the beacon code on the second page?", doc)
        assert response["grounded"]
        assert "CEDAR" in response["answer"] and "AMBER" not in response["answer"]
        assert response["citations"]
        for citation in response["citations"]:
            evidence = client.get('/api/v1/evidence/' + citation["evidence_unit_id"]).json()
            assert citation["page"] == evidence["page"] == 2
            assert citation["snippet"] in evidence["snippet"]


def test_docx_retry_replaces_stale_page_count_and_preserves_original(tmp_path):
    from app.api.dependencies import get_repository
    from app.main import app

    document = Document()
    document.add_paragraph("Opening page.")
    document.add_page_break()
    document.add_paragraph("The closing instructions use the east gate.")
    original = payload(document)
    with client_with_data_root(tmp_path) as client:
        doc = upload(client, "gate.docx", original)
        repo = app.dependency_overrides[get_repository]()
        repo.update_document(doc, pages=9)
        response = client.post(f"/api/v1/documents/{doc}/retry")
        assert response.status_code == 202, response.text
        assert repo.get_document(doc).pages == 2
        assert client.get(f"/api/v1/documents/{doc}/file").content == original
        answer = ask(client, "Quote the paragraph on the last page.", doc)
        assert answer["answer"] == "The closing instructions use the east gate."
        assert answer["citations"][0]["page"] == 2
